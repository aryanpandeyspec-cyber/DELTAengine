import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

# --- PALETTE DEFINITION (PROJECT NEUBRUTALIST THEME) ---
BG_PAGE = RGBColor(246, 245, 240)       # #F6F5F0 Warm Ivory
CHARCOAL = RGBColor(17, 24, 39)          # #111827 Dark Border & Text
CARD_WHITE = RGBColor(255, 255, 255)     # #FFFFFF
SHADOW_COLOR = RGBColor(17, 24, 39)      # Brutalist offset shadow

# Accent Colors
BLUE_ACCENT = RGBColor(37, 99, 235)      # #2563EB
BLUE_LIGHT = RGBColor(219, 234, 254)     # #DBEAFE
GREEN_ACCENT = RGBColor(16, 185, 129)    # #10B981
GREEN_LIGHT = RGBColor(220, 252, 231)    # #DCFCE7
AMBER_ACCENT = RGBColor(245, 158, 11)    # #F59E0B
AMBER_LIGHT = RGBColor(254, 243, 199)    # #FEF3C7
RED_ACCENT = RGBColor(239, 68, 68)       # #EF4444
RED_LIGHT = RGBColor(254, 226, 226)      # #FEE2E2
PURPLE_ACCENT = RGBColor(139, 92, 246)   # #8B5CF6
PURPLE_LIGHT = RGBColor(243, 232, 255)   # #F3E8FF
MUTED_TEXT = RGBColor(75, 85, 99)        # #4B5563

FONT_HEADING = 'Arial Black'
FONT_BODY = 'Segoe UI'
FONT_MONO = 'Consolas'

def set_slide_background(slide, prs):
    bg_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = BG_PAGE
    bg_shape.line.fill.background()
    return bg_shape

def add_neubrutalist_card(slide, left, top, width, height, bg_rgb=CARD_WHITE, border_rgb=CHARCOAL, shadow=True):
    if shadow:
        offset = Inches(0.06)
        s_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left + offset, top + offset, width, height)
        s_card.fill.solid()
        s_card.fill.fore_color.rgb = SHADOW_COLOR
        s_card.line.fill.background()

    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_rgb
    card.line.color.rgb = border_rgb
    card.line.width = Pt(2.2)
    return card

def add_badge(slide, left, top, width, height, text, bg_rgb, text_rgb=CHARCOAL, border_rgb=CHARCOAL):
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    badge.fill.solid()
    badge.fill.fore_color.rgb = bg_rgb
    badge.line.color.rgb = border_rgb
    badge.line.width = Pt(1.5)
    
    tf = badge.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.CENTER
    p.font.name = FONT_BODY
    p.font.bold = True
    p.font.size = Pt(10)
    p.font.color.rgb = text_rgb
    return badge

def add_header(slide, badge_text, title_text, subtitle_text, slide_num):
    # Slide Tag Badge
    add_badge(slide, Inches(0.8), Inches(0.4), Inches(3.0), Inches(0.35), badge_text, AMBER_LIGHT, CHARCOAL)
    
    # Page indicator (Total 07 slides)
    page_badge = add_badge(slide, Inches(11.3), Inches(0.4), Inches(1.2), Inches(0.35), f"0{slide_num} / 07", BLUE_LIGHT, BLUE_ACCENT)
    
    # Title
    t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.82), Inches(11.7), Inches(0.55))
    tf = t_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.name = FONT_HEADING
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = CHARCOAL

    # Subtitle
    s_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.7), Inches(0.35))
    stf = s_box.text_frame
    stf.word_wrap = True
    stf.margin_left = stf.margin_right = stf.margin_top = stf.margin_bottom = 0
    sp = stf.paragraphs[0]
    sp.text = subtitle_text
    sp.font.name = FONT_BODY
    sp.font.size = Pt(13)
    sp.font.color.rgb = MUTED_TEXT

def build_presentation():
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # =========================================================================
    # SLIDE 1: TITLE & EXECUTIVE VISION
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, prs)

    # Top Banner Badges (All HackIndia / Spark traces removed)
    add_badge(s1, Inches(0.8), Inches(0.5), Inches(3.2), Inches(0.38), "⚡ AUTONOMOUS EVENT OS", AMBER_LIGHT, CHARCOAL)
    add_badge(s1, Inches(4.15), Inches(0.5), Inches(3.2), Inches(0.38), "🏛️ ENTERPRISE ARCHITECTURE", BLUE_LIGHT, BLUE_ACCENT)

    # Left Hero Box
    left_card = add_neubrutalist_card(s1, Inches(0.8), Inches(1.15), Inches(6.8), Inches(4.7), CARD_WHITE)
    ltf = left_card.text_frame
    ltf.word_wrap = True
    ltf.margin_left = Inches(0.4)
    ltf.margin_right = Inches(0.4)
    ltf.margin_top = Inches(0.4)

    p1 = ltf.paragraphs[0]
    p1.text = "DELTA ENGINE"
    p1.font.name = FONT_HEADING
    p1.font.size = Pt(40)
    p1.font.bold = True
    p1.font.color.rgb = CHARCOAL

    p2 = ltf.add_paragraph()
    p2.text = "Autonomous, Agentic & Self-Healing Event OS"
    p2.font.name = FONT_BODY
    p2.font.size = Pt(16)
    p2.font.bold = True
    p2.font.color.rgb = BLUE_ACCENT
    p2.space_before = Pt(6)

    p3 = ltf.add_paragraph()
    p3.text = "A zero-latency operating system that autonomously resolves conference schedule conflicts, speaker delays, and venue capacity overshoots in real time without human intervention."
    p3.font.name = FONT_BODY
    p3.font.size = Pt(12)
    p3.font.color.rgb = MUTED_TEXT
    p3.space_before = Pt(12)

    # 3 Mini Value Pillars inside Left Card
    pillars = [
        ("🤖 4-Agent Groq Swarm", "Hybrid parallel-sequential LLM intelligence"),
        ("⚡ <150ms Self-Healing", "Deterministic graph heuristic solvers"),
        ("📱 Multi-Channel Dispatch", "Automated WhatsApp AI & DKIM emails")
    ]
    for title, desc in pillars:
        pp = ltf.add_paragraph()
        pp.text = f"• {title}: {desc}"
        pp.font.name = FONT_BODY
        pp.font.size = Pt(11)
        pp.font.color.rgb = CHARCOAL
        pp.space_before = Pt(8)

    # Right Hero Box (Image Presentation)
    right_card = add_neubrutalist_card(s1, Inches(7.8), Inches(1.15), Inches(4.7), Inches(4.7), CARD_WHITE)
    if os.path.exists('logo.jpeg'):
        s1.shapes.add_picture('logo.jpeg', Inches(8.0), Inches(1.35), width=Inches(4.3))
    
    # Mini stats banner beneath image on right
    rs_badge = add_badge(s1, Inches(8.0), Inches(4.75), Inches(4.3), Inches(0.85), 
                         "🌟 ZERO-LATENCY EVENT OS\nValidated with 500 High-Stress Scenarios", 
                         GREEN_LIGHT, GREEN_ACCENT)

    # Bottom Presenter Ribbon
    bot_card = add_neubrutalist_card(s1, Inches(0.8), Inches(6.1), Inches(11.7), Inches(0.9), CARD_WHITE)
    btf = bot_card.text_frame
    btf.word_wrap = True
    btf.margin_left = Inches(0.3)
    btf.margin_top = Inches(0.16)
    bp1 = btf.paragraphs[0]
    bp1.text = "TEAM DELTA  |  2nd Year CSM-C"
    bp1.font.name = FONT_BODY
    bp1.font.size = Pt(11.5)
    bp1.font.bold = True
    bp1.font.color.rgb = CHARCOAL

    bp2 = btf.add_paragraph()
    bp2.text = "25BK2A66D5 - ARYAN PANDEY   •   25BK1A66J2 - SHAIK SHAHIDUDDIN   •   25BK1A66J6 - SURYANSH SINGH"
    bp2.font.name = FONT_MONO
    bp2.font.size = Pt(11)
    bp2.font.bold = True
    bp2.font.color.rgb = BLUE_ACCENT
    bp2.space_before = Pt(2)


    # =========================================================================
    # SLIDE 2: PROBLEM VS. THE PARADIGM SHIFT
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2, prs)
    add_header(s2, "01 | THE PARADIGM SHIFT", "Why Large-Scale Events Break Down", 
               "Moving from manual coordination panic to deterministic autonomous self-healing", 2)

    # Left Card: The Status Quo (Fragile & Manual)
    c1 = add_neubrutalist_card(s2, Inches(0.8), Inches(1.9), Inches(5.65), Inches(5.1), RED_LIGHT, RED_ACCENT)
    add_badge(s2, Inches(1.1), Inches(2.15), Inches(3.2), Inches(0.38), "🚨 TRADITIONAL EVENT CHAOS", RED_ACCENT, CARD_WHITE, RED_ACCENT)
    
    if os.path.exists('doodles/15711543.png'):
        s2.shapes.add_picture('doodles/15711543.png', Inches(4.9), Inches(2.1), width=Inches(1.2))

    t1 = c1.text_frame
    t1.word_wrap = True
    t1.margin_left = Inches(0.35)
    t1.margin_right = Inches(0.35)
    t1.margin_top = Inches(0.85)

    bullets_bad = [
        ("Cascade Schedule Clashes", "A single 30-min speaker delay triggers overlapping hall conflicts across subsequent sessions."),
        ("Dangerous Hall Overflow", "Unregulated crowd surges exceed room safety capacities by up to 212% without warning."),
        ("Fragmented Phone/Chat Calls", "Coordinators, volunteers & AV staff operate in disconnected silos with no single source of truth.")
    ]
    for b_title, b_desc in bullets_bad:
        p = t1.add_paragraph()
        p.text = f"❌ {b_title}"
        p.font.name = FONT_BODY
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = CHARCOAL
        p.space_before = Pt(10)

        pd = t1.add_paragraph()
        pd.text = b_desc
        pd.font.name = FONT_BODY
        pd.font.size = Pt(11)
        pd.font.color.rgb = MUTED_TEXT
        pd.space_before = Pt(2)

    # Stat Callout Bad
    s1_box = add_neubrutalist_card(s2, Inches(1.1), Inches(5.9), Inches(5.05), Inches(0.85), CARD_WHITE, RED_ACCENT, shadow=False)
    st1 = s1_box.text_frame
    st1.margin_left = Inches(0.2)
    st1.margin_top = Inches(0.12)
    stp = st1.paragraphs[0]
    stp.text = "⏱️  45+ Minutes Average Human Conflict Resolution Time"
    stp.font.name = FONT_BODY
    stp.font.bold = True
    stp.font.size = Pt(12)
    stp.font.color.rgb = RED_ACCENT

    # Right Card: The DELTA OS (Autonomous Self-Healing)
    c2 = add_neubrutalist_card(s2, Inches(6.85), Inches(1.9), Inches(5.65), Inches(5.1), GREEN_LIGHT, GREEN_ACCENT)
    add_badge(s2, Inches(7.15), Inches(2.15), Inches(3.2), Inches(0.38), "✨ DELTA AUTONOMOUS OS", GREEN_ACCENT, CARD_WHITE, GREEN_ACCENT)

    if os.path.exists('ioncs/healing.png'):
        s2.shapes.add_picture('ioncs/healing.png', Inches(10.9), Inches(2.1), width=Inches(1.2))

    t2 = c2.text_frame
    t2.word_wrap = True
    t2.margin_left = Inches(0.35)
    t2.margin_right = Inches(0.35)
    t2.margin_top = Inches(0.85)

    bullets_good = [
        ("Autonomous Graph Reallocation", "Constraint algorithms instantly shift slots and execute clean room swaps in <150ms."),
        ("Dynamic Capacity Protection", "Monitors hall capacity matrices (250, 120, 60 pax) and balances audience distribution."),
        ("Unified AI Emergency Dispatch", "Auto-dispatches WhatsApp messages & DKIM-signed HTML emails to all affected parties.")
    ]
    for b_title, b_desc in bullets_good:
        p = t2.add_paragraph()
        p.text = f"✅ {b_title}"
        p.font.name = FONT_BODY
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = CHARCOAL
        p.space_before = Pt(10)

        pd = t2.add_paragraph()
        pd.text = b_desc
        pd.font.name = FONT_BODY
        pd.font.size = Pt(11)
        pd.font.color.rgb = MUTED_TEXT
        pd.space_before = Pt(2)

    # Stat Callout Good
    s2_box = add_neubrutalist_card(s2, Inches(7.15), Inches(5.9), Inches(5.05), Inches(0.85), CARD_WHITE, GREEN_ACCENT, shadow=False)
    st2 = s2_box.text_frame
    st2.margin_left = Inches(0.2)
    st2.margin_top = Inches(0.12)
    st2p = st2.paragraphs[0]
    st2p.text = "⚡  < 150 ms End-to-End Autonomous Resolution Latency"
    st2p.font.name = FONT_BODY
    st2p.font.bold = True
    st2p.font.size = Pt(12)
    st2p.font.color.rgb = GREEN_ACCENT


    # =========================================================================
    # SLIDE 3: SYSTEM ARCHITECTURE & AUTONOMOUS PIPELINE
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3, prs)
    add_header(s3, "02 | ARCHITECTURE & DATA FLOW", "Closed-Loop Autonomous Pipeline", 
               "From real-time disruption telemetry to instant multi-channel dispatch", 3)

    stages = [
        ("1. Telemetry Ingestion", "doodles/15818548.png", BLUE_LIGHT, BLUE_ACCENT, 
         ["Speaker delays ingested", "Hall attendance sensors", "REST API & UI triggers"]),
        ("2. Heuristic Solver", "ioncs/pipeline.png", AMBER_LIGHT, AMBER_ACCENT, 
         ["10 Bounded iterations", "Availability formulas", "Optimal slot/hall swaps"]),
        ("3. Groq Swarm AI", "ioncs/negotiate.png", PURPLE_LIGHT, PURPLE_ACCENT, 
         ["4 Parallel LLM agents", "Facility & AV checks", "Consensus logs generated"]),
        ("4. Resync & Dispatch", "ioncs/graph.png", GREEN_LIGHT, GREEN_ACCENT, 
         ["In-Memory graph rebind", "Full WebSocket broadcast", "WhatsApp & Email push"])
    ]

    col_w = Inches(2.72)
    gap = Inches(0.27)
    left_start = Inches(0.8)

    for i, (title, img_path, bg_color, accent_color, points) in enumerate(stages):
        pos_x = left_start + i * (col_w + gap)
        sc = add_neubrutalist_card(s3, pos_x, Inches(1.9), col_w, Inches(4.0), CARD_WHITE)
        
        add_badge(s3, pos_x + Inches(0.15), Inches(2.05), col_w - Inches(0.3), Inches(0.38), 
                  title, bg_color, accent_color, accent_color)

        if os.path.exists(img_path):
            s3.shapes.add_picture(img_path, pos_x + Inches(0.76), Inches(2.55), width=Inches(1.2))

        stf = sc.text_frame
        stf.word_wrap = True
        stf.margin_left = Inches(0.18)
        stf.margin_right = Inches(0.18)
        stf.margin_top = Inches(1.95)

        for pt in points:
            p = stf.add_paragraph()
            p.text = f"• {pt}"
            p.font.name = FONT_BODY
            p.font.size = Pt(10.5)
            p.font.color.rgb = CHARCOAL
            p.space_before = Pt(6)

    # Bottom Architecture Highlights
    b_bar = add_neubrutalist_card(s3, Inches(0.8), Inches(6.1), Inches(11.7), Inches(0.95), BLUE_LIGHT, BLUE_ACCENT)
    btf3 = b_bar.text_frame
    btf3.word_wrap = True
    btf3.margin_left = Inches(0.3)
    btf3.margin_top = Inches(0.15)
    
    b3_p1 = btf3.paragraphs[0]
    b3_p1.text = "KEY ARCHITECTURAL SAFEGUARDS & SPECS"
    b3_p1.font.name = FONT_HEADING
    b3_p1.font.size = Pt(11)
    b3_p1.font.bold = True
    b3_p1.font.color.rgb = BLUE_ACCENT

    b3_p2 = btf3.add_paragraph()
    b3_p2.text = "⚡ 10 Passes Max Iteration Bounded Guard  |  🕸️ In-Memory Neo4j-Style Relational Graph  |  🔄 Low-Latency WebSocket Broadcast (<50ms)"
    b3_p2.font.name = FONT_BODY
    b3_p2.font.size = Pt(11)
    b3_p2.font.bold = True
    b3_p2.font.color.rgb = CHARCOAL
    b3_p2.space_before = Pt(3)


    # =========================================================================
    # SLIDE 4: HETEROGENEOUS MULTI-AGENT AI SWARM
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4, prs)
    add_header(s4, "03 | ARTIFICIAL INTELLIGENCE", "Heterogeneous 4-Agent Groq Swarm", 
               "Hybrid parallel-sequential LLM pipeline specialized for real-time event governance", 4)

    agents = [
        ("🗣️ Liaison Agent", "llama-3.1-8b-instant", BLUE_LIGHT, BLUE_ACCENT,
         "Ingests raw attendee feedback, speaker delay telemetry, and venue anomaly flags into structured 1-sentence situational updates."),
        ("⏱️ Scheduler Agent", "llama-3.3-70b-versatile", AMBER_LIGHT, AMBER_ACCENT,
         "Executes high-order constraint logic to evaluate temporal conflicts, slot migrations, and speaker availability windows."),
        ("🏛️ Logistics Agent", "llama-3.3-70b-versatile", GREEN_LIGHT, GREEN_ACCENT,
         "Audits hall seating limits (Turing: 250, Lovelace: 120, Hopper: 60), AV technical equipment, and HVAC facilities balancing."),
        ("📢 Marketing Agent", "llama-3.1-8b-instant", PURPLE_LIGHT, PURPLE_ACCENT,
         "Drafts real-time attendee announcements, updates public visual display banners, and synchronizes iCal calendar subscriptions.")
    ]

    card_w = Inches(5.65)
    card_h = Inches(1.85)

    coords = [
        (Inches(0.8), Inches(1.9)),
        (Inches(6.85), Inches(1.9)),
        (Inches(0.8), Inches(3.95)),
        (Inches(6.85), Inches(3.95))
    ]

    for (title, model_name, bg_c, acc_c, desc), (pos_x, pos_y) in zip(agents, coords):
        ac = add_neubrutalist_card(s4, pos_x, pos_y, card_w, card_h, CARD_WHITE)
        
        add_badge(s4, pos_x + Inches(0.25), pos_y + Inches(0.18), Inches(2.2), Inches(0.32), 
                  title, bg_c, acc_c, acc_c)
        add_badge(s4, pos_x + Inches(2.6), pos_y + Inches(0.18), Inches(2.8), Inches(0.32), 
                  model_name, BG_PAGE, CHARCOAL)

        atf = ac.text_frame
        atf.word_wrap = True
        atf.margin_left = Inches(0.25)
        atf.margin_right = Inches(0.25)
        atf.margin_top = Inches(0.65)

        ap = atf.paragraphs[0]
        ap.text = desc
        ap.font.name = FONT_BODY
        ap.font.size = Pt(11)
        ap.font.color.rgb = CHARCOAL

    # Bottom Circuit Breaker Callout
    cb_card = add_neubrutalist_card(s4, Inches(0.8), Inches(6.0), Inches(11.7), Inches(1.05), AMBER_LIGHT, AMBER_ACCENT)
    cbtf = cb_card.text_frame
    cbtf.word_wrap = True
    cbtf.margin_left = Inches(0.3)
    cbtf.margin_top = Inches(0.18)

    cb_p1 = cbtf.paragraphs[0]
    cb_p1.text = "⚡ SUPER ADMIN LLM CIRCUIT BREAKER & DETERMINISTIC FALLBACK"
    cb_p1.font.name = FONT_HEADING
    cb_p1.font.size = Pt(11.5)
    cb_p1.font.bold = True
    cb_p1.font.color.rgb = AMBER_ACCENT

    cb_p2 = cbtf.add_paragraph()
    cb_p2.text = "In case of external LLM API outages or rate limits, the engine instantly engages a deterministic rule-based solver. Live conference operations never freeze or drop a single event."
    cb_p2.font.name = FONT_BODY
    cb_p2.font.size = Pt(11)
    cb_p2.font.color.rgb = CHARCOAL
    cb_p2.space_before = Pt(3)


    # =========================================================================
    # SLIDE 5: REAL-TIME OPERATIONAL DISPATCH & SECURITY
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5, prs)
    add_header(s5, "04 | DISPATCH & INFRASTRUCTURE", "Field Dispatch & Hardened Reliability", 
               "Multi-channel attendee alerts backed by enterprise-grade crash guards", 5)

    cols = [
        ("WhatsApp AI Liaison", "ioncs/live.png", GREEN_LIGHT, GREEN_ACCENT, [
            "Automated Direct Dispatch: Instant alerts sent to room leads and speakers via wa.me protocol.",
            "Volunteer Duty Registry: Direct tracking for stage leads, mic runners, and door scanners.",
            "Live Telemetry Feeds: Renders incoming delivery receipts directly inside coordinator dashboard."
        ]),
        ("Anti-Spam Supabase Email", "ioncs/DB.png", BLUE_LIGHT, BLUE_ACCENT, [
            "DKIM / SPF Signed: Clean HTML templates prevent critical updates from hitting spam folders.",
            "Audit Trail Logging: Reallocation notices persisted permanently in Supabase PostgreSQL.",
            "Resilient Offline Mode: Local caching guarantees zero notification loss during network blips."
        ]),
        ("Security & Crash Guards", "doodles/2682340.png", RED_LIGHT, RED_ACCENT, [
            "Process Crash Guards: Node uncaughtException listeners guarantee server never crashes.",
            "Express Rate Limiter: Protects all API endpoints against DoS floods and abuse.",
            "Strict File Validation: Multer memory storage enforces strict 20MB buffer limits & sanitization."
        ])
    ]

    col3_w = Inches(3.7)
    gap3 = Inches(0.3)

    for i, (title, img_p, bg_c, acc_c, points) in enumerate(cols):
        pos_x = Inches(0.8) + i * (col3_w + gap3)
        col_card = add_neubrutalist_card(s5, pos_x, Inches(1.9), col3_w, Inches(5.1), CARD_WHITE)

        add_badge(s5, pos_x + Inches(0.2), Inches(2.1), col3_w - Inches(0.4), Inches(0.38),
                  title, bg_c, acc_c, acc_c)

        if os.path.exists(img_p):
            s5.shapes.add_picture(img_p, pos_x + Inches(1.25), Inches(2.6), width=Inches(1.2))

        ctf = col_card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = Inches(0.25)
        ctf.margin_right = Inches(0.25)
        ctf.margin_top = Inches(2.05)

        for pt in points:
            title_part, desc_part = pt.split(": ")
            p = ctf.add_paragraph()
            p.text = f"• {title_part}:"
            p.font.name = FONT_BODY
            p.font.bold = True
            p.font.size = Pt(11)
            p.font.color.rgb = CHARCOAL
            p.space_before = Pt(8)

            pd = ctf.add_paragraph()
            pd.text = desc_part
            pd.font.name = FONT_BODY
            pd.font.size = Pt(10.5)
            pd.font.color.rgb = MUTED_TEXT
            pd.space_before = Pt(2)


    # =========================================================================
    # SLIDE 6: BENCHMARKS, IMPACT & PRODUCTION DEPLOYMENT
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6, prs)
    add_header(s6, "05 | BENCHMARKS & ROADMAP", "Production Impact & Live Deployment", 
               "Proven performance under extreme stress, deployed live in production", 6)

    kpis = [
        ("< 150 ms", "Autonomous Conflict Healing", GREEN_LIGHT, GREEN_ACCENT),
        ("500 / 500", "Concurrent Stress Tests Passed", BLUE_LIGHT, BLUE_ACCENT),
        ("4 Parallel LLMs", "Groq Swarm Reasoning", PURPLE_LIGHT, PURPLE_ACCENT),
        ("100% Automated", "Zero Human Latency Required", AMBER_LIGHT, AMBER_ACCENT)
    ]

    stat_w = Inches(2.72)
    stat_gap = Inches(0.27)

    for i, (metric, label, bg_c, acc_c) in enumerate(kpis):
        pos_x = Inches(0.8) + i * (stat_w + stat_gap)
        sc = add_neubrutalist_card(s6, pos_x, Inches(1.9), stat_w, Inches(1.6), bg_c, acc_c)
        stf = sc.text_frame
        stf.word_wrap = True
        stf.margin_left = Inches(0.15)
        stf.margin_right = Inches(0.15)
        stf.margin_top = Inches(0.25)

        p = stf.paragraphs[0]
        p.text = metric
        p.font.name = FONT_HEADING
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = acc_c
        p.alignment = PP_ALIGN.CENTER

        pl = stf.add_paragraph()
        pl.text = label
        pl.font.name = FONT_BODY
        pl.font.size = Pt(11)
        pl.font.bold = True
        pl.font.color.rgb = CHARCOAL
        pl.alignment = PP_ALIGN.CENTER
        pl.space_before = Pt(4)

    # Bottom Split Layout (2 Big Cards)
    # Left Card: Live Deployment & Roadmap
    left_bot = add_neubrutalist_card(s6, Inches(0.8), Inches(3.75), Inches(5.7), Inches(3.25), CARD_WHITE)
    add_badge(s6, Inches(1.05), Inches(3.95), Inches(3.4), Inches(0.35), 
              "🚀 LIVE PRODUCTION DEPLOYMENT", BLUE_LIGHT, BLUE_ACCENT, BLUE_ACCENT)

    lbtf = left_bot.text_frame
    lbtf.word_wrap = True
    lbtf.margin_left = Inches(0.3)
    lbtf.margin_right = Inches(0.3)
    lbtf.margin_top = Inches(0.7)

    lp1 = lbtf.paragraphs[0]
    lp1.text = "Live URL: https://deltaengine.vercel.app/"
    lp1.font.name = FONT_MONO
    lp1.font.size = Pt(11.5)
    lp1.font.bold = True
    lp1.font.color.rgb = BLUE_ACCENT

    roadmap_pts = [
        "Eliminates event scheduling delays and prevents room overcrowding.",
        "Deterministic fallbacks ensure zero downtime even during API surges.",
        "Next Roadmap: IoT room occupancy sensors & Apple Wallet pass sync."
    ]
    for pt in roadmap_pts:
        p = lbtf.add_paragraph()
        p.text = f"• {pt}"
        p.font.name = FONT_BODY
        p.font.size = Pt(11)
        p.font.color.rgb = CHARCOAL
        p.space_before = Pt(6)

    # Right Card: Updated TEAM DELTA Card (As requested: Roll numbers & names, 2nd Year CSM-C, clean)
    right_bot = add_neubrutalist_card(s6, Inches(6.8), Inches(3.75), Inches(5.7), Inches(3.25), CARD_WHITE)
    add_badge(s6, Inches(7.05), Inches(3.95), Inches(2.2), Inches(0.35), 
              "TEAM DELTA", AMBER_LIGHT, CHARCOAL, AMBER_ACCENT)
    add_badge(s6, Inches(9.35), Inches(3.95), Inches(2.8), Inches(0.35), 
              "2nd Year CSM-C", BLUE_LIGHT, BLUE_ACCENT, BLUE_ACCENT)

    rbtf = right_bot.text_frame
    rbtf.word_wrap = True
    rbtf.margin_left = Inches(0.35)
    rbtf.margin_right = Inches(0.35)
    rbtf.margin_top = Inches(0.85)

    team_members = [
        "25BK2A66D5 - ARYAN PANDEY",
        "25BK1A66J2 - SHAIK SHAHIDUDDIN",
        "25BK1A66J6 - SURYANSH SINGH"
    ]
    for i, member in enumerate(team_members):
        p = rbtf.add_paragraph() if i > 0 else rbtf.paragraphs[0]
        p.text = f"👤 {member}"
        p.font.name = FONT_MONO
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = CHARCOAL
        p.space_before = Pt(8) if i > 0 else Pt(0)

    rf = rbtf.add_paragraph()
    rf.text = "DELTA ENGINE  |  Autonomous Event Operating System"
    rf.font.name = FONT_BODY
    rf.font.size = Pt(10.5)
    rf.font.bold = True
    rf.font.color.rgb = MUTED_TEXT
    rf.space_before = Pt(18)


    # =========================================================================
    # SLIDE 7: THANK YOU & Q&A (PAGE 7)
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7, prs)
    add_header(s7, "06 | CONCLUSION & Q&A", "Thank You!", 
               "Autonomous & Self-Healing Event Operating System — Open for Discussion", 7)

    # Hero Thank You Card (Left)
    ty_card = add_neubrutalist_card(s7, Inches(0.8), Inches(1.9), Inches(6.8), Inches(5.1), CARD_WHITE)
    tytf = ty_card.text_frame
    tytf.word_wrap = True
    tytf.margin_left = Inches(0.4)
    tytf.margin_right = Inches(0.4)
    tytf.margin_top = Inches(0.4)

    typ1 = tytf.paragraphs[0]
    typ1.text = "THANK YOU!"
    typ1.font.name = FONT_HEADING
    typ1.font.size = Pt(44)
    typ1.font.bold = True
    typ1.font.color.rgb = CHARCOAL

    typ2 = tytf.add_paragraph()
    typ2.text = "Any Questions or Discussion?"
    typ2.font.name = FONT_BODY
    typ2.font.size = Pt(18)
    typ2.font.bold = True
    typ2.font.color.rgb = BLUE_ACCENT
    typ2.space_before = Pt(4)

    typ3 = tytf.add_paragraph()
    typ3.text = "DELTA ENGINE demonstrates how autonomous multi-agent swarms and in-memory graph solvers can completely eliminate human panic in large-scale live events."
    typ3.font.name = FONT_BODY
    typ3.font.size = Pt(12)
    typ3.font.color.rgb = MUTED_TEXT
    typ3.space_before = Pt(14)

    # 3 Summary metric pills
    summary_pts = [
        ("⚡ < 150ms Resolution", "Sub-second autonomous healing latency"),
        ("🤖 4 Groq AI Agents", "Parallel temporal, logistics, and broadcast logic"),
        ("🛡️ 100% Reliability", "500/500 high-concurrency stress test scenarios passed")
    ]
    for tag, desc in summary_pts:
        p = tytf.add_paragraph()
        p.text = f"• {tag}: {desc}"
        p.font.name = FONT_BODY
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = CHARCOAL
        p.space_before = Pt(8)

    # Right Card (Team & Deployment info)
    ty_right = add_neubrutalist_card(s7, Inches(7.8), Inches(1.9), Inches(4.7), Inches(5.1), CARD_WHITE)
    
    add_badge(s7, Inches(8.1), Inches(2.15), Inches(2.0), Inches(0.35), 
              "TEAM DELTA", AMBER_LIGHT, CHARCOAL, AMBER_ACCENT)
    add_badge(s7, Inches(10.25), Inches(2.15), Inches(2.0), Inches(0.35), 
              "2nd Year CSM-C", BLUE_LIGHT, BLUE_ACCENT, BLUE_ACCENT)

    ty_rtf = ty_right.text_frame
    ty_rtf.word_wrap = True
    ty_rtf.margin_left = Inches(0.3)
    ty_rtf.margin_right = Inches(0.3)
    ty_rtf.margin_top = Inches(0.85)

    rp1 = ty_rtf.paragraphs[0]
    rp1.text = "PROJECT AUTHORS"
    rp1.font.name = FONT_HEADING
    rp1.font.size = Pt(11)
    rp1.font.bold = True
    rp1.font.color.rgb = MUTED_TEXT

    for member in team_members:
        p = ty_rtf.add_paragraph()
        p.text = f"👤 {member}"
        p.font.name = FONT_MONO
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = CHARCOAL
        p.space_before = Pt(8)

    # Divider & Deployment link
    dp = ty_rtf.add_paragraph()
    dp.text = "───────────────"
    dp.font.name = FONT_BODY
    dp.font.size = Pt(9)
    dp.font.color.rgb = MUTED_TEXT
    dp.space_before = Pt(10)

    url_title = ty_rtf.add_paragraph()
    url_title.text = "🌐 LIVE PRODUCTION SYSTEM:"
    url_title.font.name = FONT_BODY
    url_title.font.size = Pt(10)
    url_title.font.bold = True
    url_title.font.color.rgb = CHARCOAL
    url_title.space_before = Pt(4)

    url_p = ty_rtf.add_paragraph()
    url_p.text = "https://deltaengine.vercel.app/"
    url_p.font.name = FONT_MONO
    url_p.font.size = Pt(11)
    url_p.font.bold = True
    url_p.font.color.rgb = BLUE_ACCENT
    url_p.space_before = Pt(2)

    # Save to disk
    output_filename = "DELTA_ENGINE_Presentation.pptx"
    prs.save(output_filename)
    print(f"Presentation generated successfully with 7 slides: {output_filename}")

if __name__ == '__main__':
    build_presentation()
